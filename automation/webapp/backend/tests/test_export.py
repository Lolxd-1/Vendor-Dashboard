"""Tests for app.engine.export: category mapping, row validation, workbook build.

These guard the two ways this module can silently wreck a catalogue: mapping a
menu item into the wrong SmartBiz product category, or writing to a fresh
workbook that loses the template's data validations and fails the SmartBiz
import.
"""
import io

import openpyxl
import pytest

from app.engine import export


# ---------------------------------------------------------------------------
# product_category_for
# ---------------------------------------------------------------------------

def test_ai_suggestion_wins_when_legal():
    result = export.product_category_for("Anything", "Bakery", "Other Food and Grocery")
    assert result == "Bakery"


def test_ai_suggestion_ignored_when_not_in_legal_list():
    # AI hallucinated a category that isn't one of the 15 legal strings ->
    # falls through to the keyword map on the menu category.
    result = export.product_category_for("Cold Coffee", "Hot Beverages", "Other Food and Grocery")
    assert result == "Beverages"


def test_falls_back_to_default_when_nothing_matches():
    result = export.product_category_for("Miscellaneous Stuff", None, "Other Food and Grocery")
    assert result == "Other Food and Grocery"


def test_classics_does_not_map_to_beverages_word_boundary():
    """Regression: a naive substring match sees "c-LASSI-cs" inside "Classics"."""
    result = export.product_category_for("Classics", None, "Other Food and Grocery")
    assert result != "Beverages"
    assert result == "Other Food and Grocery"


@pytest.mark.parametrize("menu_category", ["Shakes", "Cold Coffee", "Coolers"])
def test_beverage_keywords_do_map_to_beverages(menu_category):
    result = export.product_category_for(menu_category, None, "Other Food and Grocery")
    assert result == "Beverages"


def test_export_food_categories_are_all_legal_and_only_15():
    assert len(export.EXPORT_FOOD_CATEGORIES) == 15
    assert len(set(export.EXPORT_FOOD_CATEGORIES)) == 15


@pytest.mark.parametrize("menu_category,ai_suggestion", [
    ("Classics", None, ),
    ("Random Section", "Not A Real Category"),
    ("Shakes", None),
    ("", None),
    ("Anything", "Bakery"),
])
def test_product_category_for_never_escapes_legal_list(menu_category, ai_suggestion):
    result = export.product_category_for(menu_category, ai_suggestion, "Other Food and Grocery")
    assert result in export.EXPORT_FOOD_CATEGORIES


# ---------------------------------------------------------------------------
# validate()
# ---------------------------------------------------------------------------

def _row(**overrides):
    row = {
        "id": "item-1",
        "name": "Paneer Butter Masala",
        "price": 249.0,
        "category": "Main Course",
        "product_category": "Other Food and Grocery",
        "description": "Rich, creamy tomato gravy.",
        "imgbb_url": "https://i.ibb.co/example.jpg",
    }
    row.update(overrides)
    return row


def test_validate_flags_name_too_long():
    errors = export.validate([_row(name="x" * 201)])
    assert any(e.field == "name" for e in errors)


def test_validate_allows_name_at_limit():
    errors = export.validate([_row(name="x" * 200)])
    assert not any(e.field == "name" for e in errors)


def test_validate_flags_description_too_long():
    errors = export.validate([_row(description="x" * 2001)])
    assert any(e.field == "description" for e in errors)


def test_validate_allows_description_at_limit():
    errors = export.validate([_row(description="x" * 2000)])
    assert not any(e.field == "description" for e in errors)


def test_validate_flags_zero_or_negative_price():
    errors = export.validate([_row(price=0)])
    assert any(e.field == "price" for e in errors)

    errors = export.validate([_row(price=-5)])
    assert any(e.field == "price" for e in errors)


def test_validate_flags_price_over_max_mrp():
    errors = export.validate([_row(price=1_000_000.00)])
    assert any(e.field == "price" for e in errors)


def test_validate_allows_price_at_max_mrp():
    errors = export.validate([_row(price=export.MAX_MRP)])
    assert not any(e.field == "price" for e in errors)


def test_validate_flags_price_with_more_than_two_decimals():
    errors = export.validate([_row(price=10.999)])
    assert any(e.field == "price" for e in errors)


def test_validate_flags_missing_image_url():
    errors = export.validate([_row(imgbb_url=None)])
    assert any(e.field in ("imgbb_url", "image") for e in errors)


def test_validate_flags_illegal_product_category():
    # ai_suggestion not legal AND menu category doesn't match any keyword AND
    # default itself is illegal -> product_category_for returns the illegal
    # default, and validate() must catch that rather than write it silently.
    errors = export.validate([_row(
        category="Miscellaneous",
        product_category="Not A Real Category",
        default_product_category="Also Not Real",
    )])
    assert any(e.field == "product_category" for e in errors)


def test_validate_returns_empty_for_clean_row():
    errors = export.validate([_row()])
    assert errors == []


# ---------------------------------------------------------------------------
# build_workbook() against the REAL template
# ---------------------------------------------------------------------------

ROWS = [
    {
        "id": "item-1",
        "name": "Paneer Butter Masala",
        "price": 249.0,
        "category": "Main Course",
        "product_category": "Other Food and Grocery",
        "description": "Rich, creamy tomato gravy with cubes of paneer.",
        "imgbb_url": "https://i.ibb.co/abc123/paneer.jpg",
    },
    {
        "id": "item-2",
        "name": "Cold Coffee",
        "price": 99.0,
        "category": "Shakes",
        "product_category": None,
        "description": "",
        "imgbb_url": "https://i.ibb.co/abc123/coldcoffee.jpg",
    },
]


def test_build_workbook_loads_real_template():
    data = export.build_workbook(ROWS, "FOOD_AND_GROCERY")
    assert isinstance(data, bytes)
    assert len(data) > 0


def test_build_workbook_output_is_reopenable_xlsx():
    data = export.build_workbook(ROWS, "FOOD_AND_GROCERY")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert export.SHEET in wb.sheetnames


def test_build_workbook_sku_and_variant_columns_are_empty():
    data = export.build_workbook(ROWS, "FOOD_AND_GROCERY")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[export.SHEET]
    for row in range(2, 2 + len(ROWS)):
        assert ws.cell(row=row, column=1).value in (None, "")  # A: SKU ID
        assert ws.cell(row=row, column=2).value in (None, "")  # B: Variant ID


def test_build_workbook_places_fields_in_correct_columns():
    data = export.build_workbook(ROWS, "FOOD_AND_GROCERY")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb[export.SHEET]

    row = 2
    assert ws.cell(row=row, column=export.COL_NAME).value == "Paneer Butter Masala"
    assert ws.cell(row=row, column=export.COL_MRP).value == 249.0
    assert ws.cell(row=row, column=export.COL_BUSINESS_CAT).value == "FOOD_AND_GROCERY"
    assert ws.cell(row=row, column=export.COL_PRODUCT_CAT).value == "Other Food and Grocery"
    assert ws.cell(row=row, column=export.COL_DESCRIPTION).value == (
        "Rich, creamy tomato gravy with cubes of paneer."
    )
    assert ws.cell(row=row, column=export.COL_IMAGE1).value == (
        "https://i.ibb.co/abc123/paneer.jpg"
    )

    # Second row resolves its product_category via the keyword fallback
    # ("Shakes" -> Beverages) since product_category is None on this row.
    row2 = 3
    assert ws.cell(row=row2, column=export.COL_PRODUCT_CAT).value == "Beverages"


def test_build_workbook_preserves_all_sheets_and_data_validations():
    """The regression that would break every SmartBiz import: building a fresh
    workbook (instead of writing into a copy of the template) silently drops
    the template's 3 sheets and its 29 data validations."""
    data = export.build_workbook(ROWS, "FOOD_AND_GROCERY")
    wb = openpyxl.load_workbook(io.BytesIO(data))

    assert len(wb.sheetnames) == 3
    assert set(wb.sheetnames) == {
        "important_instructions", "bulk_upload_template", "DataSheet",
    }

    ws = wb[export.SHEET]
    assert len(ws.data_validations.dataValidation) == 29


def test_build_workbook_handles_empty_rows_without_error():
    data = export.build_workbook([], "FOOD_AND_GROCERY")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert export.SHEET in wb.sheetnames
    assert len(wb[export.SHEET].data_validations.dataValidation) == 29
