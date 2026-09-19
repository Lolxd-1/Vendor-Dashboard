#!/usr/bin/env python3
"""
Fill the SmartBiz (by Amazon) bulk-upload template with the generated menu.

Writes into a COPY of the real template (templates/smartbiz_template.xlsx) so
every dropdown, named range and validation rule SmartBiz expects stays intact —
building a lookalike sheet from scratch is what makes these imports fail.

Column mapping (template sheet: "bulk_upload_template")
    D  Product Name (Mandatory, <=200 chars)  <- Item Name
    E  MRP (Mandatory, number > 0)            <- Price
    G  Business Category (Mandatory, dropdown)<- --business-category
    H  Product Category (Mandatory, dropdown) <- --product-category
    I  Product Description (Optional)         <- Category column, if present
    P  Product Image1 (Optional)              <- public imgbb URL

SKU ID / Variant ID (A, B) are deliberately left blank — the template marks
them "Not to be Edited"; SmartBiz assigns them on import.

Usage
    python smartbiz_export.py                    # only items that have a hosted image
    python smartbiz_export.py --all              # include items with no image yet
    python smartbiz_export.py --product-category "Dairy"
"""

import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from dotenv import load_dotenv

import imgbb
from export_excel import load_progress, scan_images, link_for, INPUT_DIR, OUTPUT_DIR
from menu_source import pick_input_excel

load_dotenv()

BASE_DIR = Path(__file__).parent.resolve()
TEMPLATE = BASE_DIR / "templates" / "smartbiz_template.xlsx"
SHEET = "bulk_upload_template"

# 1-indexed columns in the template
COL_NAME = 4
COL_MRP = 5
COL_BUSINESS_CAT = 7
COL_PRODUCT_CAT = 8
COL_DESCRIPTION = 9
COL_IMAGE1 = 16

MAX_NAME = 200
MAX_DESC = 2000
MAX_MRP = 999999.99

DEFAULT_BUSINESS_CAT = "FOOD_AND_GROCERY"
DEFAULT_PRODUCT_CAT = "Chocolates, desserts and icecream"

# Product Category is a mandatory dropdown, and every value here must be spelled
# exactly as it appears in the template's DataSheet under FOOD_AND_GROCERY.
# A menu section called "Shakes" or "Coolers" is a drink, not a dessert, so send
# it to Beverages instead of the blanket default. Matched on the menu category
# text, longest keyword first. Anything unmatched keeps --product-category.
CATEGORY_KEYWORDS = [
    ("milkshake", "Beverages"),
    ("smoothie", "Beverages"),
    ("mocktail", "Beverages"),
    ("beverage", "Beverages"),
    ("cooler", "Beverages"),
    ("coffee", "Beverages"),
    ("shake", "Beverages"),
    ("juice", "Beverages"),
    ("lassi", "Beverages"),
    ("drink", "Beverages"),
    ("soda", "Beverages"),
    ("tea", "Beverages"),
]


def product_category_for(menu_category: str, default: str) -> str:
    """Map a menu section heading onto a valid SmartBiz product category.

    The keyword must START A WORD, not merely appear somewhere in the text:
    a plain substring test files "Classics" under Beverages, because
    "c-LASSI-cs" contains "lassi". A trailing \\w* still catches the plurals
    that real menus use ("Shakes", "Coolers", "Juices").
    """
    text = (menu_category or "").lower()
    for keyword, smartbiz_cat in CATEGORY_KEYWORDS:
        if re.search(rf"\b{re.escape(keyword)}\w*", text):
            return smartbiz_cat
    return default


def read_menu(src: Path):
    """Read the source menu Excel into dicts keyed by sheet row."""
    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    col = {}
    for i, h in enumerate(headers):
        if ("item" in h and "name" in h) or h in ("name", "item name"):
            col["name"] = i
        elif "categ" in h:
            col["category"] = i
        elif "price" in h or "mrp" in h:
            col["price"] = i
        elif "desc" in h:
            col["description"] = i
    if "name" not in col:
        raise ValueError(f"No item-name column found. Headers: {headers}")

    items = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or col["name"] >= len(row) or not row[col["name"]]:
            continue
        def get(k):
            i = col.get(k)
            return row[i] if i is not None and i < len(row) else None
        items.append({
            "row": row_idx,
            "name": str(row[col["name"]]).strip(),
            "category": str(get("category") or "").strip(),
            "price": get("price"),
            "description": str(get("description") or "").strip(),
        })
    wb.close()
    return items


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--all", action="store_true",
                    help="include items that have no hosted image yet")
    ap.add_argument("--business-category", default=DEFAULT_BUSINESS_CAT)
    ap.add_argument("--product-category", default=DEFAULT_PRODUCT_CAT)
    ap.add_argument("--no-auto-category", action="store_true",
                    help="put every row in --product-category, even drinks")
    ap.add_argument("--template", default=str(TEMPLATE))
    args = ap.parse_args()

    template = Path(args.template)
    if not template.exists():
        print(f"X SmartBiz template not found: {template}")
        sys.exit(1)

    src, why = pick_input_excel(INPUT_DIR)
    if src is None:
        print(f"X {why}")
        sys.exit(1)

    items = read_menu(src)
    progress = load_progress()
    images = scan_images()

    print(f"  menu        : {src.name}  ({len(items)} items)  [{why}]")
    print(f"  images      : {len(images)}")
    print(f"  business cat: {args.business_category}")
    print(f"  product cat : {args.product_category}")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = OUTPUT_DIR / f"smartbiz_upload_{ts}.xlsx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, out)          # keep the template's validations intact

    wb = openpyxl.load_workbook(out)
    if SHEET not in wb.sheetnames:
        print(f"X Template has no '{SHEET}' sheet (found: {wb.sheetnames})")
        sys.exit(1)
    ws = wb[SHEET]

    written = 0
    cat_counts = {}
    skipped_no_image = 0
    warnings = []
    out_row = 2

    for it in items:
        url = link_for(it["row"], progress, images)
        hosted = url.startswith("http")

        if not hosted and not args.all:
            skipped_no_image += 1
            continue

        name = it["name"]
        if len(name) > MAX_NAME:
            warnings.append(f"row {it['row']}: name truncated to {MAX_NAME} chars")
            name = name[:MAX_NAME]

        price = it["price"]
        try:
            price = float(price)
            if price <= 0 or price > MAX_MRP:
                raise ValueError
            price = round(price, 2)
        except (TypeError, ValueError):
            warnings.append(f"row {it['row']}: '{name}' has invalid MRP ({it['price']!r}) — left blank, SmartBiz will reject this row")
            price = None

        ws.cell(row=out_row, column=COL_NAME, value=name)
        ws.cell(row=out_row, column=COL_MRP, value=price)
        ws.cell(row=out_row, column=COL_BUSINESS_CAT, value=args.business_category)
        prod_cat = (args.product_category if args.no_auto_category
                    else product_category_for(it["category"], args.product_category))
        cat_counts[prod_cat] = cat_counts.get(prod_cat, 0) + 1
        ws.cell(row=out_row, column=COL_PRODUCT_CAT, value=prod_cat)
        if it["description"] or it["category"]:
            desc = (it["description"] or it["category"])[:MAX_DESC]
            ws.cell(row=out_row, column=COL_DESCRIPTION, value=desc)
        if hosted:
            ws.cell(row=out_row, column=COL_IMAGE1, value=url)
        else:
            warnings.append(f"row {it['row']}: '{name}' has no hosted image URL")

        out_row += 1
        written += 1

    wb.save(out)
    wb.close()

    print(f"\nOK  {written} products written")
    for cat, n in sorted(cat_counts.items(), key=lambda kv: -kv[1]):
        print(f"      {n:>4}  {cat}")
    if skipped_no_image:
        print(f"    {skipped_no_image} skipped (no hosted image yet — use --all to include)")
    if warnings:
        print(f"\n  {len(warnings)} warning(s):")
        for w in warnings[:12]:
            print(f"    - {w}")
        if len(warnings) > 12:
            print(f"    ... and {len(warnings) - 12} more")
    print(f"\n    {out}")


if __name__ == "__main__":
    main()
