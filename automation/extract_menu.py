#!/usr/bin/env python3
"""
STAGE 1 - Menu photographs  ->  input/menu_items.xlsx

Reads every menu photo in `input/Menu Images/` with Gemini vision and writes the
sheet the rest of the pipeline already understands:

    Item Name | Category | Price | Description | Source Menu | Image Link

`Image Link` is left empty on purpose - main.py fills it once the food images
are generated and hosted.

Built for BATCHES. Ten photos is a normal run:
  * every photo is one API call, run on a small thread pool
  * each photo's result is cached by file hash in output/menu_extract_cache.json,
    so re-running is free and adding an 11th photo costs exactly one call
  * one unreadable photo is reported and skipped - it never kills the run
  * items repeated across overlapping photos are de-duplicated

Usage
    python extract_menu.py                 # extract everything not cached
    python extract_menu.py --force         # ignore the cache, re-read every photo
    python extract_menu.py --dry-run       # print what was found, write nothing
    python extract_menu.py --dir "path"    # read photos from somewhere else
"""

import argparse
import hashlib
import json
import logging
import os
import re
import shutil
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# Windows consoles default to cp1252 and raise UnicodeEncodeError on the box
# characters and spinner glyphs Rich draws with. Force UTF-8 before Rich builds
# its Console (it samples the encoding at construction time).
for _stream in ("stdout", "stderr"):
    _s = getattr(sys, _stream, None)
    if _s is not None and hasattr(_s, "reconfigure"):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from dotenv import load_dotenv
from rich.console import Console
from rich.panel import Panel
from rich.progress import BarColumn, Progress, SpinnerColumn, TextColumn, TimeElapsedColumn
from rich.table import Table

BASE_DIR = Path(__file__).parent.resolve()
load_dotenv(BASE_DIR / ".env")

INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
MENU_DIR = INPUT_DIR / "Menu Images"
OUT_XLSX = INPUT_DIR / "menu_items.xlsx"
CACHE_FILE = OUTPUT_DIR / "menu_extract_cache.json"
BACKUP_DIR = OUTPUT_DIR / "backups"
LOG_FILE = OUTPUT_DIR / "extract_menu.log"

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_BACKUP_KEY = os.getenv("GEMINI_BACKUP_KEY", "")

# Vision model. The image *generator* (gemini-3.1-flash-lite-image) cannot read
# text back out - this is a separate, much cheaper text-out call.
VISION_MODEL = "gemini-2.5-flash"

MAX_WORKERS = 3            # text-out quota is far kinder than image-gen quota
MAX_RETRIES = 5
API_TIMEOUT = 180          # seconds; a dense board can take ~30s to read
MAX_DIM = 1600             # downscale before upload - keeps small print legible
JPEG_QUALITY = 92
MAX_OUTPUT_TOKENS = 32768  # a very dense board can emit a lot of JSON

RATE_BACKOFF_BASE = 10.0
RATE_BACKOFF_MAX = 120.0

# Bump when the prompt changes so cached reads are invalidated automatically.
PROMPT_VERSION = 6

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".jfif", ".tif", ".tiff", ".gif"}

console = Console()
logger = logging.getLogger("extract_menu")


def setup_logging():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    logger.setLevel(logging.INFO)
    if logger.handlers:
        return
    fh = logging.FileHandler(LOG_FILE, encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    logger.addHandler(fh)


# ============================================================
# PROMPT
# ============================================================

RESPONSE_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "is_menu": {"type": "BOOLEAN"},
        "restaurant_name": {"type": "STRING"},
        "currency": {"type": "STRING"},
        "items": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "item_name": {"type": "STRING"},
                    "category": {"type": "STRING"},
                    "price": {"type": "NUMBER"},
                    "description": {"type": "STRING"},
                },
                "required": ["item_name", "category", "price"],
            },
        },
    },
    "required": ["is_menu", "items"],
}

EXTRACT_PROMPT = """You are reading a photograph of a restaurant menu (a menu board, printed card, or leaflet).

Extract EVERY purchasable item that has a price printed next to it.

Rules:
1. ONE ROW PER PRICE. If an item shows several prices under column headings
   (Large/Small, Double layer/Single layer, Half/Full, Box of 4/Box of 6),
   emit one row per price and put the variant in the name:
   "Choco Brownie Sundae (Large)", "Almond Brownie Cake (Double Layer)".
2. If a single name contains a slash for interchangeable options
   ("Honey Butter/Maple Butter", "Iced Tea (Lemon/Peach)"), that is ONE item -
   emit one row and keep the name exactly as printed.
3. category = the most specific section heading printed above the item
   ("Classics", "Shakes", "Waffle Cakes", "Coolers"). Never invent a category.
   If the section has a parent heading, use the specific child heading.
   Do NOT repeat the category inside the item name. Under a "Mini Waff-wich
   Combos" heading, "Box of 4 / Chocoholic's" becomes item_name
   "Chocoholic's (Box of 4)" with category "Mini Waff-wich Combos" - never
   "Mini Waff-wich Combos Box of 4 Chocoholic's".
4. price = the number only. No currency symbol, no commas.
5. description = the small print printed under that item, if any. Otherwise "".
6. Copy names EXACTLY as printed, including accents. Fix only obvious OCR
   damage, never re-word or expand a name.
7. A priced upgrade or add-on IS an item ("Add an ice-cream scoop - Rs 30",
   "Turn any Classic No-Maida - Rs 30"). Category it exactly like anything
   else: use its printed section heading if it sits under one (a toppings
   block headed "Extra Goodness" stays "Extra Goodness"). Only use "Add-ons"
   for a standalone priced banner that has no section heading at all.
8. IGNORE anything with no price of its own: allergen and copyright notices,
   GST lines, phone numbers, addresses, social handles, opening hours,
   decorative slogans, and any adjacent poster, screen or signboard that is not
   part of this menu itself.
9. Do NOT guess. If a price or a name is genuinely unreadable in the photo,
   skip that item rather than inventing a value.
10. Set is_menu=false and items=[] if this photograph is not a menu at all.

Return strict JSON matching the schema."""


# ============================================================
# CACHE
# ============================================================

_cache_lock = threading.Lock()


def load_cache() -> dict:
    if not CACHE_FILE.exists():
        return {}
    try:
        return json.loads(CACHE_FILE.read_text(encoding="utf-8"))
    except Exception as e:
        logger.warning(f"cache unreadable ({e}) - starting empty")
        return {}


def save_cache(cache: dict):
    CACHE_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = CACHE_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(CACHE_FILE)


def file_key(path: Path) -> str:
    """Content hash + prompt version - rename-proof, edit-aware."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return f"v{PROMPT_VERSION}_{h.hexdigest()[:32]}"


# ============================================================
# GEMINI VISION
# ============================================================

_RATE_MARKERS = ("429", "resource_exhausted", "resource has been exhausted",
                 "quota", "rate limit", "too many requests")
_AUTH_MARKERS = ("401", "403", "unauthenticated", "permission_denied",
                 "api key not valid", "invalid api key", "billing")


def _is_rate_limit(msg: str) -> bool:
    m = msg.lower()
    return any(k in m for k in _RATE_MARKERS)


def _is_auth_error(msg: str) -> bool:
    m = msg.lower()
    return any(k in m for k in _AUTH_MARKERS)


def _finish_reason(resp) -> str:
    try:
        return str(resp.candidates[0].finish_reason)
    except Exception:
        return "?"


class MenuReader:
    """One Vertex express-mode client, shared across worker threads."""

    def __init__(self, api_key: str, backup_key: str = ""):
        self._key = api_key
        self._backup = backup_key
        self._client = None
        self._client_key = None
        self._lock = threading.Lock()
        self.calls = 0

    def _client_for(self):
        from google import genai
        from google.genai import types
        with self._lock:
            if self._client is None or self._client_key != self._key:
                self._client = genai.Client(
                    vertexai=True,
                    api_key=self._key,
                    http_options=types.HttpOptions(timeout=API_TIMEOUT * 1000),
                )
                self._client_key = self._key
            return self._client

    def _switch_to_backup(self) -> bool:
        with self._lock:
            if self._backup and self._key != self._backup:
                self._key = self._backup
                logger.warning("switched to BACKUP key")
                return True
        return False

    @staticmethod
    def prepare(path: Path) -> bytes:
        """Downscale to MAX_DIM and re-encode as JPEG."""
        im = Image.open(path)
        # honour EXIF rotation - a sideways photo reads badly
        try:
            from PIL import ImageOps
            im = ImageOps.exif_transpose(im)
        except Exception:
            pass
        im = im.convert("RGB")
        if max(im.size) > MAX_DIM:
            im.thumbnail((MAX_DIM, MAX_DIM), Image.LANCZOS)
        buf = BytesIO()
        im.save(buf, format="JPEG", quality=JPEG_QUALITY)
        return buf.getvalue()

    def read(self, path: Path) -> dict:
        """Read one menu photo. Returns the parsed dict. Raises on final failure."""
        from google.genai import types

        jpeg = self.prepare(path)
        part = types.Part.from_bytes(data=jpeg, mime_type="image/jpeg")
        config = types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=RESPONSE_SCHEMA,
            temperature=0.0,
            max_output_tokens=MAX_OUTPUT_TOKENS,
        )

        last_err = "unknown error"
        for attempt in range(MAX_RETRIES + 1):
            with self._lock:
                self.calls += 1
                n = self.calls
            logger.info(f"  API call #{n} '{path.name}' [try {attempt + 1}/{MAX_RETRIES + 1}]")
            try:
                resp = self._client_for().models.generate_content(
                    model=VISION_MODEL, contents=[part, EXTRACT_PROMPT], config=config,
                )
            except Exception as e:
                msg = str(e)
                last_err = f"{type(e).__name__}: {msg[:250]}"
                if _is_rate_limit(msg):
                    wait = min(RATE_BACKOFF_BASE * (2 ** attempt), RATE_BACKOFF_MAX)
                    logger.warning(f"  [429] {path.name} - sleeping {wait:.0f}s")
                    time.sleep(wait)
                    continue
                if _is_auth_error(msg):
                    logger.error(f"  [AUTH] {msg[:250]}")
                    if self._switch_to_backup():
                        continue
                    raise RuntimeError(f"AUTH/BILLING FAILURE: {msg[:250]}")
                wait = min(4 * (2 ** attempt), 45)
                logger.warning(f"  [ERR] {last_err} - retry in {wait:.0f}s")
                time.sleep(wait)
                continue

            text = (resp.text or "").strip()
            if not text:
                last_err = f"empty response (finish_reason={_finish_reason(resp)})"
                logger.warning(f"  {path.name}: {last_err}")
                time.sleep(3)
                continue
            try:
                data = json.loads(text)
            except json.JSONDecodeError as e:
                # Nearly always a token-truncated response.
                last_err = f"invalid JSON ({e}); {len(text)} chars returned"
                logger.warning(f"  {path.name}: {last_err}")
                time.sleep(3)
                continue

            if not isinstance(data, dict) or "items" not in data:
                last_err = f"unexpected JSON shape: {str(data)[:150]}"
                time.sleep(3)
                continue
            return data

        raise RuntimeError(last_err)


# ============================================================
# NORMALISE + MERGE
# ============================================================

_WS = re.compile(r"\s+")


def clean_text(v) -> str:
    return _WS.sub(" ", str(v or "")).strip()


def norm_key(name: str, category: str) -> str:
    """Dedupe key: case/space/punctuation-insensitive name + category."""
    def squash(s: str) -> str:
        return re.sub(r"[^a-z0-9]+", "", s.lower())
    return f"{squash(name)}|{squash(category)}"


def parse_price(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        p = float(v)
    else:
        s = re.sub(r"[^\d.]", "", str(v))
        if not s:
            return None
        try:
            p = float(s)
        except ValueError:
            return None
    if p <= 0 or p > 999999:
        return None
    return round(p, 2)


def normalise_items(raw: List[dict], source: str) -> Tuple[List[dict], List[str]]:
    """Clean one photo's items. Returns (items, warnings)."""
    items, warns = [], []
    for r in raw or []:
        name = clean_text(r.get("item_name"))
        if not name:
            continue
        if len(name) > 200:
            warns.append(f"{source}: '{name[:40]}...' name truncated to 200 chars")
            name = name[:200]
        price = parse_price(r.get("price"))
        if price is None:
            warns.append(f"{source}: '{name}' has no usable price ({r.get('price')!r}) - dropped")
            continue
        items.append({
            "item_name": name,
            "category": clean_text(r.get("category")),
            "price": price,
            "description": clean_text(r.get("description"))[:2000],
            "source": source,
        })
    return items, warns


def merge(per_photo: List[Tuple[str, List[dict]]]) -> Tuple[List[dict], List[str]]:
    """Concatenate photos in order, dropping repeats across overlapping shots."""
    seen: Dict[str, dict] = {}
    merged, notes = [], []
    for source, items in per_photo:
        for it in items:
            k = norm_key(it["item_name"], it["category"])
            prev = seen.get(k)
            if prev is None:
                seen[k] = it
                merged.append(it)
                continue
            if abs(prev["price"] - it["price"]) < 0.01:
                notes.append(
                    f"duplicate skipped: '{it['item_name']}' ({it['category']}) "
                    f"also in {prev['source']}")
            else:
                # Same name+category, different price -> genuinely different rows
                # (or a misread). Keep both and flag it for a human to glance at.
                notes.append(
                    f"PRICE CONFLICT: '{it['item_name']}' ({it['category']}) = "
                    f"{prev['price']:g} in {prev['source']} but {it['price']:g} "
                    f"in {source} - both kept")
                merged.append(it)
    return merged, notes


# ============================================================
# EXCEL OUT
# ============================================================

HEADERS = ["Item Name", "Category", "Price", "Description", "Source Menu", "Image Link"]
COL_WIDTHS = [42, 22, 10, 52, 34, 56]


def write_excel(items: List[dict], path: Path):
    if path.exists():
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = BACKUP_DIR / f"{path.stem}_{stamp}{path.suffix}"
        shutil.copy2(path, backup)
        console.print(f"[dim]  previous sheet backed up -> {backup.name}[/]")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Menu"

    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS[c - 1]
    ws.freeze_panes = "A2"

    for r, it in enumerate(items, start=2):
        ws.cell(row=r, column=1, value=it["item_name"])
        ws.cell(row=r, column=2, value=it["category"])
        ws.cell(row=r, column=3, value=it["price"])
        ws.cell(row=r, column=4, value=it["description"])
        ws.cell(row=r, column=5, value=it["source"])
        # column 6 (Image Link) is filled later by the image pipeline

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(items) + 1)}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


# ============================================================
# DISCOVERY
# ============================================================

def find_photos(menu_dir: Path) -> Tuple[List[Path], List[Path]]:
    """Return (usable photos, files we had to skip)."""
    if not menu_dir.exists():
        return [], []
    photos, skipped = [], []
    for f in sorted(menu_dir.rglob("*")):
        if not f.is_file() or f.name.startswith("~$") or f.name.startswith("."):
            continue
        if f.suffix.lower() in IMAGE_EXTS:
            if f.stat().st_size > 0:
                photos.append(f)
            else:
                skipped.append(f)
        else:
            skipped.append(f)
    return photos, skipped


# ============================================================
# MAIN
# ============================================================

def run(menu_dir: Path, force: bool, dry_run: bool, workers: int) -> int:
    setup_logging()
    console.print(Panel("[bold white]STAGE 1 - READ MENU PHOTOS -> EXCEL[/]",
                        border_style="bright_cyan", padding=(1, 4)))

    photos, skipped = find_photos(menu_dir)
    if not photos:
        console.print(f"[bold red]X No menu photos found in {menu_dir}[/]")
        console.print(f"  Drop your menu photographs ({', '.join(sorted(IMAGE_EXTS))}) there and re-run.")
        return 1
    if not GEMINI_API_KEY:
        console.print("[bold red]X GEMINI_API_KEY is not set in .env[/]")
        return 1

    cache = load_cache()
    todo = [p for p in photos if force or file_key(p) not in cache]

    info = Table.grid(padding=(0, 2))
    info.add_row("Folder", f"[cyan]{menu_dir}[/]")
    info.add_row("Photos", f"[bold]{len(photos)}[/]")
    info.add_row("Cached", f"[green]{len(photos) - len(todo)}[/]")
    info.add_row("To read", f"[yellow]{len(todo)}[/]")
    info.add_row("Model", f"[cyan]{VISION_MODEL}[/]")
    info.add_row("Workers", f"[bold]{min(workers, max(1, len(todo)))}[/]")
    console.print(info)
    if skipped:
        console.print(f"[yellow]! {len(skipped)} non-image file(s) ignored: "
                      f"{', '.join(s.name for s in skipped[:4])}"
                      f"{' ...' if len(skipped) > 4 else ''}[/]")
    console.print()

    reader = MenuReader(GEMINI_API_KEY, GEMINI_BACKUP_KEY)
    failures: List[Tuple[Path, str]] = []
    fatal: Optional[str] = None

    if todo:
        with Progress(SpinnerColumn(), TextColumn("[progress.description]{task.description}"),
                      BarColumn(bar_width=36), TextColumn("({task.completed}/{task.total})"),
                      TimeElapsedColumn(), console=console) as prog:
            task = prog.add_task("Reading menus", total=len(todo))

            def work(p: Path):
                return p, reader.read(p)

            with ThreadPoolExecutor(max_workers=max(1, min(workers, len(todo)))) as pool:
                futures = {pool.submit(work, p): p for p in todo}
                for fut in as_completed(futures):
                    p = futures[fut]
                    try:
                        _, data = fut.result()
                        n = len(data.get("items") or [])
                        if not data.get("is_menu", True) and n == 0:
                            failures.append((p, "not recognised as a menu"))
                        else:
                            with _cache_lock:
                                cache[file_key(p)] = {
                                    "file": p.name,
                                    "read_at": datetime.now().isoformat(timespec="seconds"),
                                    "restaurant_name": data.get("restaurant_name") or "",
                                    "currency": data.get("currency") or "",
                                    "items": data.get("items") or [],
                                }
                                save_cache(cache)      # crash-safe: written per photo
                            logger.info(f"  {p.name}: {n} items")
                    except RuntimeError as e:
                        if "AUTH/BILLING FAILURE" in str(e):
                            fatal = str(e)
                        failures.append((p, str(e)[:200]))
                    except Exception as e:
                        failures.append((p, f"{type(e).__name__}: {str(e)[:180]}"))
                    prog.advance(task)

    if fatal:
        console.print(f"\n[bold red]X {fatal}[/]")
        console.print("  Nothing else will succeed until the key or billing is fixed.")
        return 1

    # ---- assemble from cache, in photo order ----
    per_photo, warns = [], []
    for p in photos:
        entry = cache.get(file_key(p))
        if not entry:
            continue
        items, w = normalise_items(entry.get("items"), p.name)
        per_photo.append((p.name, items))
        warns.extend(w)

    merged, notes = merge(per_photo)

    # ---- report ----
    console.print()
    per = Table(title="Items per photo", title_style="bold", header_style="bold cyan")
    per.add_column("Menu photo", overflow="fold")
    per.add_column("Items", justify="right")
    for name, items in per_photo:
        per.add_row(name, str(len(items)))
    per.add_row("[bold]TOTAL (deduped)[/]", f"[bold]{len(merged)}[/]")
    console.print(per)

    cats: Dict[str, int] = {}
    for it in merged:
        key = it["category"] or "(none)"
        cats[key] = cats.get(key, 0) + 1
    if cats:
        cat_tbl = Table(title="Categories", title_style="bold", header_style="bold cyan")
        cat_tbl.add_column("Category")
        cat_tbl.add_column("Items", justify="right")
        for c, n in sorted(cats.items(), key=lambda kv: -kv[1]):
            cat_tbl.add_row(c, str(n))
        console.print(cat_tbl)

    for w in warns:
        console.print(f"[yellow]! {w}[/]")
    for n in notes:
        style = "red" if n.startswith("PRICE CONFLICT") else "dim"
        console.print(f"[{style}]- {n}[/]")
    if failures:
        console.print(f"\n[bold red]{len(failures)} photo(s) could not be read:[/]")
        for p, err in failures:
            console.print(f"  * {p.name}: [dim]{err}[/]")
        console.print("[yellow]  Re-run to retry just those - everything else is cached.[/]")

    if not merged:
        console.print("\n[bold red]X No items extracted - nothing written.[/]")
        return 1

    if dry_run:
        console.print("\n[yellow]--dry-run: no Excel written.[/]")
        return 0

    write_excel(merged, OUT_XLSX)
    console.print(Panel(
        f"[green bold]{len(merged)}[/] items -> [bold]{OUT_XLSX}[/]\n"
        f"[dim]Check the prices, then run:[/] [bold]python main.py[/]",
        title="[bold]Stage 1 complete[/]", border_style="green"))
    return 1 if failures else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", default=str(MENU_DIR), help="folder holding the menu photographs")
    ap.add_argument("--force", action="store_true", help="ignore the cache and re-read every photo")
    ap.add_argument("--dry-run", action="store_true", help="print the result, write no Excel")
    ap.add_argument("--workers", type=int, default=MAX_WORKERS,
                    help=f"parallel reads (default {MAX_WORKERS})")
    args = ap.parse_args()
    try:
        sys.exit(run(Path(args.dir), args.force, args.dry_run, args.workers))
    except KeyboardInterrupt:
        console.print("\n[yellow]Interrupted - finished photos are cached, just run again.[/]")
        sys.exit(130)


if __name__ == "__main__":
    main()
